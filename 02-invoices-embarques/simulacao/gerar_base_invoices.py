"""
Gera a base fictícia de invoices e embarques de exportação (ciclo completo, deadlines e estadia).
Uso: python gerar_base_invoices.py saida/Base_Invoices.xlsx saida/fInvoices.csv

Projeto de portfólio com dados 100% fictícios.
"""
import numpy as np, pandas as pd, math, os, sys
from datetime import date, timedelta
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

rng = np.random.default_rng(77)
START = date(2025,10,1); END = date(2026,9,30)
AUTOM = date(2026,3,1)   # automação documental (fictícia) reduz atrasos de documento

ORIG = {'Paranaguá':(0.45,2),'Itajaí':(0.30,2),'Santos':(0.25,3)}
DEST = [('Xangai','China','Ásia',0.14),('Qingdao','China','Ásia',0.08),('Tóquio','Japão','Ásia',0.08),('Yokohama','Japão','Ásia',0.05),
        ('Jeddah','Arábia Saudita','Oriente Médio',0.12),('Jebel Ali','Emirados Árabes Unidos','Oriente Médio',0.12),
        ('Hong Kong','Hong Kong','Ásia',0.08),('Manzanillo','México','América do Norte',0.06),('Durban','África do Sul','África',0.06),
        ('San Antonio','Chile','América do Sul',0.05),('Busan','Coreia do Sul','Ásia',0.06),('Roterdã','Países Baixos','Europa',0.05),
        ('Alexandria','Egito','África',0.05)]
ARM = [('Atlântico Lines','ATLU'),('Mar Azul Shipping','MAZU'),('Oceano Global','OGLU'),('Pacífico Container Line','PCLU'),
       ('Horizonte Marítima','HMZU'),('Nórdica Lines','NRDU')]
TRANSP = [('TR01','Rota Sul Transportes',720),('TR02','TransOeste Logística',800),('TR03','Cerrado Cargas',760),('TR04','Litoral Express',900),
          ('TR05','Serra Verde Transportes',840),('TR06','Meridional Fretes',780),('TR07','Planalto Rodoviário',880),('TR08','Norte Sul Cargas',820)]
PRODUTOS = [('Peito de frango congelado',0.20),('Coxa e sobrecoxa congeladas',0.25),('Asa de frango congelada',0.15),('Pés de frango congelados',0.12),
            ('Miúdos congelados',0.10),('Cortes variados congelados',0.18)]
USD_T = {'Peito de frango congelado':2300,'Coxa e sobrecoxa congeladas':1900,'Asa de frango congelada':2600,'Pés de frango congelados':1700,
         'Miúdos congelados':1500,'Cortes variados congelados':2000}
CAUSAS = ['Sem causa','Aguardando documento liberativo','Alteração de PCP / saldo de produto','Aguardando laudo laboratorial','Programação de embarque incorreta','Problema com veículo/transportador']
P_PRE  = [0.39,0.30,0.13,0.10,0.06,0.02]
P_POS  = [0.66,0.08,0.12,0.09,0.03,0.02]
SEAS = {10:1.0,11:1.05,12:1.1,1:0.9,2:0.85,3:1.0,4:1.0,5:1.05,6:1.0,7:1.05,8:1.1,9:1.1}

CV = {c:i for i,c in enumerate('ABCDEFGHIJKLMNOPQRSTUVWXYZ')}
VAL = {}
v = 10
for ch in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ':
    if v % 11 == 0: v += 1
    VAL[ch] = v; v += 1
def check_digit(code10):
    s = sum(VAL[c]*(2**i) if c.isalpha() else int(c)*(2**i) for i,c in enumerate(code10))
    return (s % 11) % 10
used_cont = set()
def container(prefix):
    while True:
        serial = '%06d' % int(rng.integers(100000, 999999))
        if prefix+serial in used_cont: continue
        used_cont.add(prefix+serial)
        return f"{prefix}{serial}-{check_digit(prefix+serial)}"
# placas por transportador
def placa():
    # padrão Mercosul: 3 letras + 1 dígito + 1 letra + 2 dígitos (ex.: BRA2E19)
    L = ''.join(rng.choice(list('ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 3))
    return f"{L}{rng.integers(0,10)}{rng.choice(list('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))}{rng.integers(0,10)}{rng.integers(0,10)}"
_used_pl = set()
def placa_unica():
    while True:
        p = placa()
        if p not in _used_pl:
            _used_pl.add(p); return p
PLACAS = {t[0]:[placa_unica() for _ in range(90)] for t in TRANSP}
def nxt(d):
    return d + timedelta(1) if d.weekday() == 6 else d

rows = []; seq = 0
d = START
while d <= END:
    if d.weekday() < 5:
        n = rng.poisson(6.0*SEAS[d.month])
        for _ in range(n):
            seq += 1
            orig = str(rng.choice(list(ORIG), p=[v[0] for v in ORIG.values()])); tr = ORIG[orig][1]
            dest = DEST[int(rng.choice(len(DEST), p=np.array([x[3] for x in DEST])/sum(x[3] for x in DEST)))]
            arm = ARM[int(rng.integers(0, len(ARM)))]
            prod = PRODUTOS[int(rng.choice(len(PRODUTOS), p=[x[1] for x in PRODUTOS]))][0]
            peso = int(rng.integers(24000, 27500))
            usd = int(peso/1000*USD_T[prod]*rng.uniform(0.95, 1.08))
            L = int(rng.integers(8, 14))
            ind = d
            dl_porto = ind + timedelta(L); dl_carga = dl_porto - timedelta(tr); dl_draft = dl_porto - timedelta(3)
            cancel = rng.random() < 0.02
            tp = TRANSP[int(rng.integers(0, len(TRANSP)))]; tp2 = tp
            same = rng.random() < 0.55
            if not same: tp2 = TRANSP[int(rng.integers(0, len(TRANSP)))]
            pl1 = str(rng.choice(PLACAS[tp[0]]))
            pl2 = pl1 if same and tp2 == tp else str(rng.choice(PLACAS[tp2[0]]))
            if same: tp2 = tp
            coleta = nxt(ind + timedelta(int(rng.integers(1, 4))))
            chegada = nxt(coleta + timedelta(int(rng.integers(1, 3))))
            causa = str(rng.choice(CAUSAS, p=(P_PRE if ind < AUTOM else P_POS)))
            dd = {'Aguardando documento liberativo': int(rng.integers(1,3) if ind >= AUTOM else rng.integers(1,4)),
                  'Alteração de PCP / saldo de produto': int(rng.integers(1,4)), 'Aguardando laudo laboratorial': int(rng.integers(2,5)),
                  'Programação de embarque incorreta': int(rng.integers(1,4)), 'Problema com veículo/transportador': int(rng.integers(1,3))}.get(causa, 0)
            carreg = chegada + timedelta((1 if rng.random() < 0.10 else 0) + (dd if causa in ('Alteração de PCP / saldo de produto','Programação de embarque incorreta','Problema com veículo/transportador') else 0))
            carreg = nxt(carreg)
            fat = nxt(carreg + timedelta((1 if rng.random() < 0.10 else 0) + (dd if causa == 'Aguardando laudo laboratorial' else 0)))
            doc = nxt(fat + timedelta((1 if rng.random() < 0.30 else 0) + (dd if causa == 'Aguardando documento liberativo' else 0)))
            lib = nxt(doc + timedelta(1 if rng.random() < 0.20 else 0))
            draft = nxt(doc + timedelta(int(rng.integers(0, 2))))
            porto = nxt(lib + timedelta(tr + (1 if rng.random() < 0.15 else 0)))
            dias = max(0, (lib - chegada).days - 1)
            diaria = round(tp[2]*rng.uniform(0.95, 1.2), -1)
            valor = float(dias*diaria)
            rows.append(dict(
                numero_invoice=f"INV-{ind.year}-{seq:05d}", data_indicacao=ind, porto_origem=orig, numero_container=container(arm[1]),
                data_coleta_container=coleta, armador=arm[0], porto_destino=dest[0], pais_destino=dest[1],
                deadline_draft=dl_draft, deadline_carga=dl_carga, deadline_porto=dl_porto, status_carregamento=None,
                data_carregamento=carreg, data_faturamento=fat, data_documentacao=doc, data_liberacao=lib,
                placa_veiculo_1=pl1, transportador_1=tp[1], data_chegada_veiculo=chegada, placa_veiculo_2=pl2, transportador_2=tp2[1],
                valor_estadia=valor, produto=prod, peso_kg=peso, valor_invoice_usd=usd, data_envio_draft=draft, data_chegada_porto=porto,
                dias_estadia=dias, causa_estadia=(causa if (dias > 0 and causa != 'Sem causa') else ('Sem causa registrada' if dias > 0 else None)),
                _cancel=cancel, _causa_real=causa))
    d += timedelta(1)

df = pd.DataFrame(rows)
datecols = ['data_coleta_container','data_carregamento','data_faturamento','data_documentacao','data_liberacao','data_chegada_veiculo','data_envio_draft','data_chegada_porto']
alldates = ['data_indicacao','deadline_draft','deadline_carga','deadline_porto'] + datecols
for c in alldates: df[c] = pd.to_datetime(df[c])
ENDT = pd.Timestamp(END)
# cancelados
cm = df._cancel
keep_coleta = cm & (rng.random(len(df)) < 0.5)
for c in ['data_carregamento','data_faturamento','data_documentacao','data_liberacao','data_chegada_veiculo','data_envio_draft','data_chegada_porto']:
    df.loc[cm, c] = pd.NaT
df.loc[cm & ~keep_coleta, 'data_coleta_container'] = pd.NaT
df.loc[cm, ['placa_veiculo_1','transportador_1','placa_veiculo_2','transportador_2']] = None
df.loc[cm, ['valor_estadia','dias_estadia']] = [0.0, 0]
df.loc[cm, 'causa_estadia'] = None
# truncamento por data final
for c in datecols: df.loc[df[c] > ENDT, c] = pd.NaT
df.loc[df.data_liberacao.isna(), ['valor_estadia','dias_estadia']] = [0.0, 0]
df.loc[df.data_liberacao.isna(), 'causa_estadia'] = None
df.loc[df.data_coleta_container.isna() & ~cm, ['placa_veiculo_1','transportador_1','placa_veiculo_2','transportador_2']] = None
def status(r):
    if r._cancel: return 'Cancelado'
    if pd.notna(r.data_carregamento):
        if r._causa_real == 'Aguardando laudo laboratorial' and pd.isna(r.data_faturamento): return 'Aguardando laudo'
        return 'Carregado'
    if pd.notna(r.data_chegada_veiculo): return 'Aguardando carregamento'
    if pd.notna(r.data_coleta_container): return 'Aguardando chegada do veículo'
    return 'Aguardando coleta do contêiner'
df['status_carregamento'] = df.apply(status, axis=1)
df['dias_estadia'] = df.dias_estadia.astype(int)
df['valor_estadia'] = df.valor_estadia.round(2)
df = df.drop(columns=['_cancel','_causa_real'])
df['causa_estadia'] = df.causa_estadia
ORDER_COLS = ['numero_invoice','data_indicacao','porto_origem','numero_container','data_coleta_container','armador','porto_destino','pais_destino',
 'deadline_draft','deadline_carga','deadline_porto','status_carregamento','data_carregamento','data_faturamento','data_documentacao','data_liberacao',
 'placa_veiculo_1','transportador_1','data_chegada_veiculo','placa_veiculo_2','transportador_2','valor_estadia',
 'produto','peso_kg','valor_invoice_usd','data_envio_draft','data_chegada_porto','dias_estadia','causa_estadia']
df = df[ORDER_COLS]

# dimensões
dTransp = pd.DataFrame(TRANSP, columns=['transportador_id','transportador','diaria_estadia_ref'])
dArm = pd.DataFrame([(a[0], a[1]) for a in ARM], columns=['armador','prefixo_container'])
dDest = pd.DataFrame([(x[0],x[1],x[2]) for x in DEST], columns=['porto_destino','pais_destino','regiao'])
dPorto = pd.DataFrame([(k, v[1]) for k,v in ORIG.items()], columns=['porto_origem','dias_transito_planta_porto'])
cal = pd.DataFrame({'data': pd.date_range(START, END)})
MESES = ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
cal['ano'] = cal.data.dt.year; cal['mes_num'] = cal.data.dt.month; cal['mes_nome'] = cal.mes_num.map(lambda m: MESES[m-1])
cal['ano_mes'] = cal.data.dt.strftime('%Y-%m'); cal['semana_ano'] = cal.data.dt.isocalendar().week.astype(int)

# ---- gabarito (por mês da data_indicacao, excluindo cancelados nos prazos)
df['ym'] = df.data_indicacao.dt.strftime('%Y-%m')
ok = df[df.status_carregamento != 'Cancelado']
done = ok[ok.data_liberacao.notna()]
g = {}
g['Invoices indicadas'] = df.groupby('ym').size()
g['% cancelamentos'] = (df.status_carregamento == 'Cancelado').groupby(df.ym).mean()
dd = lambda a,b: (done[a]-done[b]).dt.days
g['Lead time indicação → carregamento (dias)'] = dd('data_carregamento','data_indicacao').groupby(done.ym).mean()
g['Carregamento → faturamento (dias)'] = dd('data_faturamento','data_carregamento').groupby(done.ym).mean()
g['Faturamento → documentação (dias)'] = dd('data_documentacao','data_faturamento').groupby(done.ym).mean()
g['Documentação → liberação (dias)'] = dd('data_liberacao','data_documentacao').groupby(done.ym).mean()
g['Lead time total indicação → liberação (dias)'] = dd('data_liberacao','data_indicacao').groupby(done.ym).mean()
g['% liberadas dentro do deadline_carga'] = (done.data_liberacao <= done.deadline_carga).groupby(done.ym).mean()
g['% chegada ao porto dentro do deadline_porto'] = (done.data_chegada_porto.dropna().le(done.deadline_porto.loc[done.data_chegada_porto.notna()])).groupby(done.ym.loc[done.data_chegada_porto.notna()]).mean()
dr = done[done.data_envio_draft.notna()]
g['% draft enviado dentro do deadline_draft'] = (dr.data_envio_draft <= dr.deadline_draft).groupby(dr.ym).mean()
g['Valor total de estadia (R$)'] = df.groupby('ym').valor_estadia.sum()
g['% invoices com estadia'] = (done.dias_estadia > 0).groupby(done.ym).mean()
g['Estadia média por invoice com estadia (R$)'] = done[done.dias_estadia > 0].groupby('ym').valor_estadia.mean()
g['Dias médios de estadia (invoices com estadia)'] = done[done.dias_estadia > 0].groupby('ym').dias_estadia.mean()
g['% trocas de contêiner (mesmo veículo traz e leva)'] = (ok.placa_veiculo_1 == ok.placa_veiculo_2).groupby(ok.ym).mean()
G = pd.DataFrame(g).T
G.columns.name = None
IND = [
('Invoices indicadas','int','Contagem de invoices por data de indicação.','COUNTROWS(fInvoices)'),
('% cancelamentos','pct','Invoices canceladas ÷ total indicado.','DIVIDE(CALCULATE(COUNTROWS(fInvoices), fInvoices[status_carregamento] = "Cancelado"), COUNTROWS(fInvoices))'),
('Lead time indicação → carregamento (dias)','dec1','Média de (data_carregamento − data_indicacao) das invoices liberadas, sem canceladas.','AVERAGEX(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado"), DATEDIFF(fInvoices[data_indicacao], fInvoices[data_carregamento], DAY))'),
('Carregamento → faturamento (dias)','dec1','Média de (data_faturamento − data_carregamento). Reflete cargas aguardando laudo.','AVERAGEX(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado"), DATEDIFF(fInvoices[data_carregamento], fInvoices[data_faturamento], DAY))'),
('Faturamento → documentação (dias)','dec1','Média de (data_documentacao − data_faturamento).','AVERAGEX(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado"), DATEDIFF(fInvoices[data_faturamento], fInvoices[data_documentacao], DAY))'),
('Documentação → liberação (dias)','dec1','Média de (data_liberacao − data_documentacao).','AVERAGEX(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado"), DATEDIFF(fInvoices[data_documentacao], fInvoices[data_liberacao], DAY))'),
('Lead time total indicação → liberação (dias)','dec1','Média de (data_liberacao − data_indicacao).','AVERAGEX(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado"), DATEDIFF(fInvoices[data_indicacao], fInvoices[data_liberacao], DAY))'),
('% liberadas dentro do deadline_carga','pct','Invoices com data_liberacao ≤ deadline_carga ÷ invoices liberadas.','DIVIDE(COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado" && fInvoices[data_liberacao] <= fInvoices[deadline_carga])), COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado")))'),
('% chegada ao porto dentro do deadline_porto','pct','Invoices com data_chegada_porto ≤ deadline_porto ÷ invoices com chegada registrada (as ainda em trânsito saem do cálculo).','DIVIDE(COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_chegada_porto]) && fInvoices[data_chegada_porto] <= fInvoices[deadline_porto])), COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_chegada_porto]))))'),
('% draft enviado dentro do deadline_draft','pct','Invoices com data_envio_draft ≤ deadline_draft ÷ invoices com draft enviado.','DIVIDE(COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_envio_draft]) && fInvoices[data_envio_draft] <= fInvoices[deadline_draft])), COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_envio_draft]))))'),
('Valor total de estadia (R$)','brl','Soma de valor_estadia.','SUM(fInvoices[valor_estadia])'),
('% invoices com estadia','pct','Invoices liberadas com dias_estadia > 0 ÷ invoices liberadas.','DIVIDE(COUNTROWS(FILTER(fInvoices, fInvoices[dias_estadia] > 0 && NOT ISBLANK(fInvoices[data_liberacao]))), COUNTROWS(FILTER(fInvoices, NOT ISBLANK(fInvoices[data_liberacao]) && fInvoices[status_carregamento] <> "Cancelado")))'),
('Estadia média por invoice com estadia (R$)','brl','Valor de estadia ÷ invoices com estadia.','DIVIDE(SUM(fInvoices[valor_estadia]), COUNTROWS(FILTER(fInvoices, fInvoices[dias_estadia] > 0)))'),
('Dias médios de estadia (invoices com estadia)','dec1','Média de dias_estadia nas invoices com estadia.','AVERAGEX(FILTER(fInvoices, fInvoices[dias_estadia] > 0), fInvoices[dias_estadia])'),
('% trocas de contêiner (mesmo veículo traz e leva)','pct','Invoices em que placa_veiculo_1 = placa_veiculo_2 ÷ invoices não canceladas.','DIVIDE(COUNTROWS(FILTER(fInvoices, fInvoices[placa_veiculo_1] = fInvoices[placa_veiculo_2] && fInvoices[status_carregamento] <> "Cancelado")), COUNTROWS(FILTER(fInvoices, fInvoices[status_carregamento] <> "Cancelado")))'),
]
FMT = {'pct':'0.0%','int':'#,##0','brl':'"R$" #,##0','dec1':'0.0','dec2':'0.00'}
# quebras para análise (tabelas extras no gabarito)
bycause = done[done.dias_estadia > 0].groupby('causa_estadia').agg(invoices=('numero_invoice','count'), dias=('dias_estadia','sum'), valor_estadia=('valor_estadia','sum')).sort_values('valor_estadia', ascending=False).reset_index()
bytransp = done.groupby('transportador_1').agg(invoices=('numero_invoice','count'), com_estadia=('dias_estadia', lambda s: (s>0).sum()), valor_estadia=('valor_estadia','sum')).reset_index()
bytransp['estadia_media_por_invoice'] = (bytransp.valor_estadia/bytransp.invoices).round(2)
bypre = done.assign(periodo=np.where(done.data_indicacao < pd.Timestamp(AUTOM), 'Antes (out/25–fev/26)', 'Depois (mar/26–set/26)')).groupby('periodo').agg(invoices=('numero_invoice','count'), pct_estadia=('dias_estadia', lambda s: (s>0).mean()), valor_total=('valor_estadia','sum'), media_por_invoice=('valor_estadia','mean')).reset_index()
df = df.drop(columns='ym')
print(df.shape); print(G.round(2).to_string()); print(bycause); print(bytransp); print(bypre)
print(df.status_carregamento.value_counts()); print(df.isna().sum()[df.isna().sum()>0])
print(df.head(3).T)

# ---------------- escrever
OUT = sys.argv[1]; CSVOUT = sys.argv[2]
HF = PatternFill('solid', fgColor='1F3864'); HFont = Font(bold=True, color='FFFFFF', name='Arial', size=10); BF = Font(name='Arial', size=10)
def hdr(ws, row=1):
    for c in ws[row]:
        if c.value is not None: c.fill = HF; c.font = HFont; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
def tbl(ws, name, d):
    t = Table(displayName=name, ref=f"A1:{get_column_letter(d.shape[1])}{d.shape[0]+1}")
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(t)
leia = [
('BASE FICTÍCIA DE INVOICES E EMBARQUES DE EXPORTAÇÃO',''),('',''),
('O que é','Base 100% fictícia: ~1.500 invoices (1 contêiner por invoice) indicadas entre 01/10/2025 e 30/09/2026, com o ciclo completo: indicação, coleta do contêiner, chegada do veículo, carregamento, faturamento, documentação, liberação, draft e chegada ao porto. Nenhuma empresa, transportador, armador ou placa é real.'),
('Colunas pedidas','As 22 primeiras colunas de fInvoices seguem exatamente a lista que você pediu, na mesma ordem.'),
('Colunas extras (ao final)','produto, peso_kg, valor_invoice_usd, data_envio_draft, data_chegada_porto, dias_estadia, causa_estadia. Foram adicionadas porque os deadlines só viram indicador quando comparados com uma data real (draft enviado, chegada ao porto) e porque a estadia precisa de causa para análise.'),
('',''),
('PREMISSAS (confirme se batem com o seu processo)',''),
('Veículo 1 e 2','Você listou duas vezes "Placa veículo / Transportador". Interpretei assim: veículo 1 traz o contêiner vazio à planta (data_chegada_veiculo); veículo 2 leva o contêiner carregado ao porto. Em ~55% dos casos é o mesmo veículo (troca de contêiner: traz o vazio e leva o cheio). Se no seu caso for outra coisa (ex.: cavalo e carreta), me diga e eu ajusto.'),
('Placas','Todas no padrão Mercosul (3 letras, 1 dígito, 1 letra, 2 dígitos). Cada transportador tem 90 placas diferentes, sem repetição entre transportadores; uma mesma placa reaparece ao longo do ano, como na vida real.'),
('Número do contêiner','Formato ISO 6346 correto: 4 letras (3 do armador + U), 6 dígitos e 1 dígito verificador após o hífen (ex.: ABCU123456-7). O exemplo do pedido tinha 5 dígitos antes do hífen; o padrão real tem 6.'),
('Deadlines','deadline_porto = limite de entrada do contêiner no terminal; deadline_carga = limite para a carga sair da planta (deadline_porto menos o trânsito planta→porto: Paranaguá 2d, Itajaí 2d, Santos 3d); deadline_draft = 3 dias antes do deadline_porto.'),
('Estadia','dias_estadia = (data_liberacao − data_chegada_veiculo) em dias, menos 1 dia de franquia. valor_estadia = dias × diária do transportador (R$ 720 a R$ 900 de referência, com variação).'),
('Causas de estadia','Aguardando documento liberativo, alteração de PCP/saldo de produto, aguardando laudo laboratorial, programação de embarque incorreta e problema com veículo/transportador. Estadias sem causa preenchida aparecem como "Sem causa registrada" (de propósito, para treinar tratamento de dados).'),
('Automação documental','A partir de 01/03/2026 a simulação reduz os atrasos por documento liberativo, como se uma automação tivesse sido implantada. É uma narrativa fictícia para o dashboard mostrar antes e depois; não é dado real de nenhuma empresa.'),
('Cargas aguardando laudo','Seguem a lógica de buffer: carregam, mas só são faturadas depois da liberação do laudo (data_faturamento atrasa).'),
('Cancelamentos e em andamento','~2% de invoices canceladas (sem datas posteriores) e invoices recentes ainda em andamento (campos futuros em branco e status correspondente).'),
('',''),
('COMO USAR',''),
('Power BI','Importe a pasta de trabalho (ou o CSV fInvoices.csv). Relacione dCalendario[data] com fInvoices[data_indicacao] (ativo) e crie relacionamentos inativos para as demais datas, usando USERELATIONSHIP quando precisar. Dimensões: transportador_1 ↔ dTransportador[transportador], armador ↔ dArmador[armador], porto_destino ↔ dDestino[porto_destino], porto_origem ↔ dPorto[porto_origem].'),
('Gabarito','A aba Gabarito_Mensal traz os indicadores por mês da data de indicação, calculados em Python. Abaixo dele ficam quebras por causa, por transportador e antes/depois de mar/2026.'),
]
w = pd.ExcelWriter(OUT, engine='openpyxl', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy')
pd.DataFrame(leia, columns=['a','b']).to_excel(w, sheet_name='LeiaMe', index=False, header=False)
wl = w.sheets['LeiaMe']; wl.column_dimensions['A'].width = 34; wl.column_dimensions['B'].width = 140
for row in wl.iter_rows():
    for c in row: c.font = BF; c.alignment = Alignment(wrap_text=True, vertical='top')
wl['A1'].font = Font(bold=True, size=14, name='Arial')
for r in range(1, wl.max_row+1):
    a = wl.cell(row=r, column=1).value
    if a in ('PREMISSAS (confirme se batem com o seu processo)','COMO USAR'): wl.cell(row=r, column=1).font = Font(bold=True, name='Arial', size=11, color='1F3864')
    elif a and wl.cell(row=r, column=2).value: wl.cell(row=r, column=1).font = Font(bold=True, name='Arial', size=10)
# Indicadores_DAX
pd.DataFrame([dict(indicador=n, definicao=de, medida_DAX=f"[{n}] = {dax}") for n,f,de,dax in IND]).to_excel(w, sheet_name='Indicadores_DAX', index=False)
wi = w.sheets['Indicadores_DAX']; hdr(wi)
for r in range(2, wi.max_row+1):
    for ci in range(1,4):
        c = wi.cell(row=r, column=ci); c.alignment = Alignment(wrap_text=True, vertical='top'); c.font = Font(name='Consolas', size=9) if ci==3 else BF
for col,wd in zip('ABC',[44,60,130]): wi.column_dimensions[col].width = wd
wi.freeze_panes = 'A2'
# Gabarito
gm = G.copy(); gm.insert(0,'indicador',gm.index)
gm.to_excel(w, sheet_name='Gabarito_Mensal', index=False)
wg = w.sheets['Gabarito_Mensal']; hdr(wg)
for r in range(2, len(gm)+2):
    f = FMT[IND[r-2][1]]
    wg.cell(row=r, column=1).font = BF
    for c in range(2, 2+G.shape[1]): wg.cell(row=r, column=c).number_format = f; wg.cell(row=r, column=c).font = BF
wg.column_dimensions['A'].width = 52
for c in range(2, 2+G.shape[1]): wg.column_dimensions[get_column_letter(c)].width = 13
wg.freeze_panes = 'B2'
r0 = len(gm)+4
def put(title, d, r, fm=None):
    wg.cell(row=r, column=1, value=title).font = Font(bold=True, name='Arial', size=11, color='1F3864')
    for j,c in enumerate(d.columns):
        x = wg.cell(row=r+1, column=1+j, value=c); x.fill = HF; x.font = HFont
    for i,row in enumerate(d.itertuples(index=False)):
        for j,v in enumerate(row):
            x = wg.cell(row=r+2+i, column=1+j, value=(v.item() if hasattr(v,'item') else v)); x.font = BF
            if fm and d.columns[j] in fm: x.number_format = fm[d.columns[j]]
    return r+3+len(d)
r1 = put('Estadia por causa (invoices liberadas com estadia)', bycause, r0, {'valor_estadia':'"R$" #,##0'})
r2 = put('Estadia por transportador (veículo 1)', bytransp, r1, {'valor_estadia':'"R$" #,##0','estadia_media_por_invoice':'"R$" #,##0'})
put('Antes e depois da automação documental (01/03/2026)', bypre, r2, {'pct_estadia':'0.0%','valor_total':'"R$" #,##0','media_por_invoice':'"R$" #,##0'})
for name, d in [('fInvoices',df),('dTransportador',dTransp),('dArmador',dArm),('dDestino',dDest),('dPorto',dPorto),('dCalendario',cal)]:
    d.to_excel(w, sheet_name=name, index=False); ws = w.sheets[name]; hdr(ws); ws.freeze_panes = 'A2'
    for i,c in enumerate(d.columns, start=1):
        wd = max(len(str(c)), 10)
        if str(d[c].dtype).startswith('datetime'): wd = 12
        elif d[c].dtype == object: wd = min(max(wd, int(d[c].astype(str).str.len().max())+2), 36)
        ws.column_dimensions[get_column_letter(i)].width = wd + 2
    tbl(ws, name, d)
    if name == 'fInvoices':
        for r in range(2, len(d)+2): ws.cell(row=r, column=22).number_format = '#,##0.00'
w.book._sheets = [w.book[n] for n in ['LeiaMe','Indicadores_DAX','Gabarito_Mensal','fInvoices','dTransportador','dArmador','dDestino','dPorto','dCalendario']]
w.book.active = 0
w.close()
c = df.copy()
for col in c.columns:
    if str(c[col].dtype).startswith('datetime'): c[col] = c[col].dt.strftime('%Y-%m-%d')
c.to_csv(CSVOUT, index=False, encoding='utf-8')
